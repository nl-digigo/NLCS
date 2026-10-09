"""Verplaats in de WERKVERSIE de SBS-rijen naar OPENVERHARDING_STRAATBAKSTEEN_...
(SBS = straatbaksteen; 5.2 koppelt sobject SVH-SBS -> object OPENVERHARDING,
omschrijving OPENVERHARDING_STRAATBAKSTEEN). Naam hernoemen + segment-kolommen
opnieuw splitsen; B/kleur blijven ongewijzigd.
"""
import os
import re
import openpyxl

FOLDER = "ontwikkeling/classificatie/fase 1 Inhoud classificatie VH"
WV = os.path.join(FOLDER, "Classificatie-NLCS-VH-inhoud-werkversie.xlsx")
SHEET = "objecten-concept-5.1-VH"


def strip_prefix(s):
    s = str(s).strip()
    while True:
        m = re.match(r"^(B|V|N|O)-(.+)$", s)
        if m:
            s = m.group(2); continue
        m = re.match(r"^[AS][A-Z]{2}-(.+)$", s)
        if m:
            s = m.group(1); continue
        return s


try:
    wb = openpyxl.load_workbook(WV)
except PermissionError:
    raise SystemExit("Werkversie is open/gelockt (sluit 'm in Excel).")
ws = wb[SHEET]
hdr = [c.value for c in ws[1]]
col = {v: i + 1 for i, v in enumerate(hdr) if v}
CA = 1
SEG = [col["object"], col["subobject01"], col["subobject02"],
       col["subobject03"], col["subobject04"], col["subobject05"]]

renamed = 0
samples = []
for ri in range(2, ws.max_row + 1):
    a = ws.cell(ri, CA).value
    if not a or not str(a).strip():
        continue
    name = strip_prefix(str(a))
    parts = name.split("_")
    if parts[0].upper() != "SBS":
        continue
    new_parts = ["OPENVERHARDING", "STRAATBAKSTEEN"] + parts[1:]
    new_name = "_".join(new_parts)
    ws.cell(ri, CA).value = new_name
    sp = new_name.split("_")
    if len(sp) > len(SEG):
        sp = sp[:len(SEG) - 1] + ["_".join(sp[len(SEG) - 1:])]
    for i, c in enumerate(SEG):
        ws.cell(ri, c).value = sp[i] if i < len(sp) else None
    renamed += 1
    if len(samples) < 10:
        samples.append((name, new_name))

wb.save(WV)
print("hernoemd (SBS -> OPENVERHARDING_STRAATBAKSTEEN):", renamed)
for old, new in samples:
    print(f"   {old}  ->  {new}")
