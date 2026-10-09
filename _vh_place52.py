"""Plaats 5.2-objecten/symbolen in de WERKVERSIE onder hun hoofdobject, bepaald via
de 'sobject'-kolom (zoekfilter) in objecten-5-2-VH.csv:
  sobject (symboolnaam) -> object (hoofdobject).
Werkwijze:
 - verwijder eerder toegevoegde 5.2-rijen (oranje gemarkeerd);
 - voor elk 5.2-symbool: pak finalCleanName, zoek de langste sobject-prefix-match
   -> hoofdobject; staat de naam daar nog niet onder, plak het hoofdobject ervoor;
 - voeg alleen toe wat (genormaliseerd) nog niet in de lijst staat;
 - B = Fysiek, hoofdgroep = VH, omschrijving-cel oranje als markering 'uit 5.2'.
"""
import os
import re
import csv
import openpyxl
from openpyxl.styles import PatternFill

FOLDER = "ontwikkeling/classificatie/fase 1 Inhoud classificatie VH"
WV = os.path.join(FOLDER, "Classificatie-NLCS-VH-inhoud-werkversie.xlsx")
SHEET = "objecten-concept-5.1-VH"
MARK_RGB = "FFE0B3"
MARK_FILL = PatternFill(fgColor=MARK_RGB, fill_type="solid")


def strip_prefix(s):
    s = str(s).strip()
    while True:
        m = re.match(r"^(B|V|N|O)-(.+)$", s)
        if m:
            s = m.group(2); continue
        m = re.match(r"^[AS][A-Z]{2}-(.+)$", s)
        if m:
            s = m.group(1); continue
        return s


def normkey(s):
    s = re.sub(r"-(SO|D)$", "", strip_prefix(s), flags=re.I)
    s = re.sub(r"_KL$", "", s, flags=re.I)
    s = re.sub(r"(?<=\d) +(?=\d)", "X", s)
    return s.strip().upper()


def readrows(path):
    with open(path, encoding="utf-8-sig") as f:
        head = f.readline()
        d = ";" if head.count(";") > head.count(",") else ","
        f.seek(0)
        return list(csv.DictReader(f, delimiter=d))


OBJ = readrows("tabellen/publicatie/objectentabellen/5-2/objecten-5-2-VH.csv")
SYM = readrows("tabellen/publicatie/symbolentabellen/5-2/symbolen-5-2-VH.csv")

# sobject (zonder prefix, upper) -> meest voorkomende object (hoofdobject)
from collections import Counter, defaultdict
sob_cnt = defaultdict(Counter)
for r in OBJ:
    so = (r.get("sobject") or "").strip()
    if so:
        sob_cnt[strip_prefix(so).upper()][(r.get("object") or "").strip().upper()] += 1
sob_map = {k: c.most_common(1)[0][0] for k, c in sob_cnt.items() if c}
sob_keys = sorted(sob_map, key=len, reverse=True)


def parent_hoofdobject(cleanname):
    u = cleanname.upper()
    for k in sob_keys:
        if u == k or u.startswith(k + "_"):
            return sob_map[k]
    return None


try:
    wb = openpyxl.load_workbook(WV)
except PermissionError:
    raise SystemExit("Werkversie is open/gelockt (sluit 'm in Excel).")
ws = wb[SHEET]
hdr = [c.value for c in ws[1]]
col = {v: i + 1 for i, v in enumerate(hdr) if v}
CA = 1
CB = col["Fysiek object of ruimtelijk gebied? "]
CH = col["hoofdgroep"]
COBJ = col["object"]
CST = col.get("status")
CDI = col.get("discipline")
SEG = [col["object"], col["subobject01"], col["subobject02"],
       col["subobject03"], col["subobject04"], col["subobject05"]]


def is_marked(ri):
    f = ws.cell(ri, CA).fill
    return bool(f and f.patternType and str(f.fgColor.rgb or "").endswith(MARK_RGB))


# 1) eerder toegevoegde (oranje) 5.2-rijen verwijderen
marked = [ri for ri in range(2, ws.max_row + 1) if is_marked(ri)]
removed = len(marked)
if marked:
    ws.delete_rows(min(marked), max(marked) - min(marked) + 1)

# 2) bekende hoofdobjecten + aanwezige namen
known_hoofd = set()
present = set()
for ri in range(2, ws.max_row + 1):
    a = ws.cell(ri, CA).value
    if not a or not str(a).strip():
        continue
    present.add(normkey(a))
    o = ws.cell(ri, COBJ).value
    if o and str(o).strip():
        known_hoofd.add(str(o).strip().upper())
    known_hoofd.add(strip_prefix(str(a)).split("_")[0].upper())

# 3) kandidaten
cand = {}


def consider(name0, is_symbol):
    name = re.sub(r"-(SO|D)$", "", strip_prefix(str(name0)), flags=re.I).strip()
    if not name:
        return
    first = name.split("_")[0].upper()
    if is_symbol and first not in known_hoofd:
        parent = parent_hoofdobject(name)
        if parent and not name.upper().startswith(parent + "_"):
            name = parent + "_" + name
    k = normkey(name)
    if k not in present and k not in cand:
        cand[k] = name


for r in OBJ:
    x = r.get("omschrijving")
    if x and x.strip():
        consider(x, False)
for r in SYM:
    fcn = r.get("finalCleanName") or r.get("symbool")
    if fcn and fcn.strip():
        consider(fcn, True)

# 4) toevoegen
ri = ws.max_row + 1
for key in sorted(cand):
    name = cand[key]
    ws.cell(ri, CA).value = name
    ws.cell(ri, CA).fill = MARK_FILL
    ws.cell(ri, CB).value = "Fysiek"
    ws.cell(ri, CH).value = "VH"
    if CST:
        ws.cell(ri, CST).value = "*"
    if CDI:
        ws.cell(ri, CDI).value = "**"
    parts = name.split("_")
    if len(parts) > len(SEG):
        parts = parts[:len(SEG) - 1] + ["_".join(parts[len(SEG) - 1:])]
    for i, c in enumerate(SEG):
        ws.cell(ri, c).value = parts[i] if i < len(parts) else None
    ri += 1

wb.save(WV)
print("verwijderd (oude 5.2-rijen):", removed)
print("toegevoegd:", len(cand))
newg = Counter(cand[k].split("_")[0].upper() for k in cand
               if cand[k].split("_")[0].upper() not in known_hoofd)
print("evt. nieuwe groep-prefixes:", dict(newg))
