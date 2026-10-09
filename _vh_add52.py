"""Voeg aan de WERKVERSIE de 5.2-objecten/symbolen toe die nog niet in de lijst
staan (schijnverschillen genegeerd; 5.2-symbolen die al als deel van een object
in de lijst staan worden overgeslagen). Toegevoegde rijen:
  - omschrijving = opgeschoonde naam (SVH-/fase-prefix en -SO/-D eraf);
  - gesplitst op '_' naar object/subobject01..05;
  - hoofdgroep = VH; kolom B = 'uit 5.2 (classificeren)';
  - omschrijving-cel oranje gekleurd als markering.
Idempotent: al toegevoegde namen worden niet nogmaals toegevoegd.
"""
import os
import re
import csv
import openpyxl
from openpyxl.styles import PatternFill

FOLDER = "ontwikkeling/classificatie/fase 1 Inhoud classificatie VH"
WV = os.path.join(FOLDER, "Classificatie-NLCS-VH-inhoud-werkversie.xlsx")
SHEET = "objecten-concept-5.1-VH"
FLAG = "uit 5.2 (classificeren)"
MARK_FILL = PatternFill(fgColor="FFE0B3", fill_type="solid")


def strip_prefix(s):
    s = str(s).strip()
    while True:
        m = re.match(r"^(B|V|N|O)-(.+)$", s)
        if m:
            s = m.group(2); continue
        m = re.match(r"^[AS][A-Z]{2}-(.+)$", s)
        if m:
            s = m.group(1); continue
        return s


def clean_name(s):
    return re.sub(r"-(SO|D)$", "", strip_prefix(s), flags=re.I).strip()


def normkey(s):
    s = clean_name(s)
    s = re.sub(r"_KL$", "", s, flags=re.I)
    s = re.sub(r"(?<=\d) +(?=\d)", "X", s)
    return s.strip().upper()


def readcol(path, names, filt=None):
    out = []
    with open(path, encoding="utf-8-sig") as f:
        head = f.readline()
        d = ";" if head.count(";") > head.count(",") else ","
        f.seek(0)
        r = csv.DictReader(f, delimiter=d)
        c = next((x for x in r.fieldnames if x and x.strip().lower() in names),
                 r.fieldnames[0])
        for row in r:
            if filt and not filt(row):
                continue
            v = row.get(c)
            if v and str(v).strip():
                out.append(str(v).strip())
    return out


try:
    wb = openpyxl.load_workbook(WV)
except PermissionError:
    raise SystemExit("Werkversie is open/gelockt (sluit 'm in Excel).")
ws = wb[SHEET]
hdr = [c.value for c in ws[1]]
col = {v: i + 1 for i, v in enumerate(hdr) if v}
CA = 1
CB = col["Fysiek object of ruimtelijk gebied? "]
CH = col["hoofdgroep"]
CST = col.get("status")
CDI = col.get("discipline")
SEG = [col["object"], col["subobject01"], col["subobject02"],
       col["subobject03"], col["subobject04"], col["subobject05"]]

present = set()
for ri in range(2, ws.max_row + 1):
    a = ws.cell(ri, CA).value
    if a and str(a).strip():
        present.add(normkey(a))

cand = {}
for x in readcol("tabellen/publicatie/objectentabellen/5-2/objecten-5-2-VH.csv",
                 {"omschrijving"}):
    cand.setdefault(normkey(x), (x, "object"))
for x in readcol("tabellen/publicatie/symbolentabellen/5-2/symbolen-5-2-VH.csv",
                 {"symbool"}):
    cand.setdefault(normkey(x), (x, "symbool"))


def suffix_present(k):
    tail = "_" + k
    return any(p == k or p.endswith(tail) for p in present)


to_add = sorted((k, v[0], v[1]) for k, v in cand.items()
                if k not in present and not suffix_present(k))

ri = ws.max_row + 1
added = 0
for key, raw, bron in to_add:
    name = clean_name(raw)
    ws.cell(ri, CA).value = name
    ws.cell(ri, CA).fill = MARK_FILL
    ws.cell(ri, CB).value = FLAG
    ws.cell(ri, CH).value = "VH"
    if CST:
        ws.cell(ri, CST).value = "*"
    if CDI:
        ws.cell(ri, CDI).value = "**"
    parts = name.split("_")
    if len(parts) > len(SEG):
        parts = parts[:len(SEG) - 1] + ["_".join(parts[len(SEG) - 1:])]
    for i, c in enumerate(SEG):
        ws.cell(ri, c).value = parts[i] if i < len(parts) else None
    ri += 1
    added += 1

wb.save(WV)
print("toegevoegd:", added)
