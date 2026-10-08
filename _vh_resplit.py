"""Herstel in de WERKVERSIE de segment-kolommen (I/K/M/O/Q/S) voor rijen waar een
segment-cel een (niet-werkende) FORMULE bevat, of waar kolom M nog een
'_'-gecombineerde waarde heeft: splits kolom A (objectnaam) op '_' en vul de
segment-kolommen opnieuw. Underscore-markerkolommen blijven ongemoeid.
"""
import os
import openpyxl

FOLDER = "ontwikkeling/classificatie/fase 1 Inhoud classificatie VH"
DST = os.path.join(FOLDER, "Classificatie-NLCS-VH-inhoud-werkversie.xlsx")
SHEET = "objecten-concept-5.1-VH"

try:
    wb = openpyxl.load_workbook(DST)          # data_only=False -> formules zichtbaar
except PermissionError:
    raise SystemExit("Werkversie is open/gelockt (sluit 'm in Excel) en kan niet "
                     "worden bijgewerkt.")
ws = wb[SHEET]
hdr = [c.value for c in ws[1]]
col = {v: i + 1 for i, v in enumerate(hdr) if v}
CA = 1
CM = col["subobject02"]
SEG = [col["object"], col["subobject01"], col["subobject02"],
       col["subobject03"], col["subobject04"], col["subobject05"]]


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


fixed_formula = 0
fixed_combined = 0
samples = []
for ri in range(2, ws.max_row + 1):
    a = ws.cell(ri, CA).value
    if not a or not isinstance(a, str):
        continue
    segvals = [ws.cell(ri, c).value for c in SEG]
    has_formula = any(is_formula(v) for v in segvals)
    m = ws.cell(ri, CM).value
    combined = isinstance(m, str) and "_" in m
    if not (has_formula or combined):
        continue
    parts = a.split("_")
    if len(parts) > len(SEG):
        parts = parts[:len(SEG) - 1] + ["_".join(parts[len(SEG) - 1:])]
    for i, c in enumerate(SEG):
        ws.cell(ri, c).value = parts[i] if i < len(parts) else None
    if has_formula:
        fixed_formula += 1
    else:
        fixed_combined += 1
    if len(samples) < 10:
        samples.append((ri, a, parts))

wb.save(DST)
print("rijen met formule hersteld:", fixed_formula)
print("rijen met '_'-combinatie hersteld:", fixed_combined)
print("--- voorbeelden ---")
for ri, a, parts in samples:
    print(f"   rij {ri}: {a}  ->  {parts}")
